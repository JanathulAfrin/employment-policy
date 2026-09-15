import pytest
from fastapi.testclient import TestClient
import io
import fitz
from docx import Document as DocxDocument

from app.main import app
from app.services.extractor import extract_pdf_bytes, extract_docx_bytes
from app.services.chunker import chunk_document_data

client = TestClient(app)

def create_sample_pdf_bytes():
    """Generates a valid 2-page PDF file in memory."""
    doc = fitz.open()
    
    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Company Leave Policy 2026\nEmployees are entitled to 20 days paid leave annually.")
    
    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Remote Work Policy\nEmployees may work remotely up to two days per week with manager approval.")
    
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes():
    """Generates a valid DOCX file with headings and paragraphs in memory."""
    doc = DocxDocument()
    doc.add_heading("Code of Conduct", level=1)
    doc.add_paragraph("All employees must maintain professional behavior at all times.")
    
    doc.add_heading("Confidentiality Agreement", level=2)
    doc.add_paragraph("Company trade secrets and proprietary code must remain strictly confidential.")
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# --- Unit Tests ---

def test_extract_pdf_metadata():
    pdf_bytes = create_sample_pdf_bytes()
    extracted = extract_pdf_bytes(pdf_bytes)
    assert len(extracted) == 2
    assert extracted[0]["page_number"] == 1
    assert "Leave Policy" in extracted[0]["text"]
    assert extracted[1]["page_number"] == 2
    assert "Remote Work Policy" in extracted[1]["text"]


def test_extract_docx_headings():
    docx_bytes = create_sample_docx_bytes()
    extracted = extract_docx_bytes(docx_bytes)
    assert len(extracted) >= 2
    sections = [s["section_heading"] for s in extracted]
    assert "Code of Conduct" in sections
    assert "Confidentiality Agreement" in sections


def test_chunk_metadata_retention():
    pdf_bytes = create_sample_pdf_bytes()
    extracted = extract_pdf_bytes(pdf_bytes)
    chunks = chunk_document_data("doc_test_123", extracted, "pdf")
    
    assert len(chunks) >= 2
    for chunk in chunks:
        assert chunk["document_id"] == "doc_test_123"
        assert "chunk_id" in chunk
        assert chunk["page_number"] in [1, 2]
        assert chunk["text"] != ""


# --- API Endpoint Integration Tests ---

def test_api_health():
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json()["status"] == "ok"


def test_upload_valid_pdf():
    pdf_bytes = create_sample_pdf_bytes()
    response = client.post(
        "/api/documents/upload",
        files={"file": ("leave_policy.pdf", pdf_bytes, "application/pdf")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["message"] == "Policy uploaded and processed successfully."
    assert data["document"]["filename"] == "leave_policy.pdf"
    assert data["document"]["file_type"] == "pdf"
    assert data["document"]["page_count"] == 2
    assert data["document"]["chunk_count"] >= 2


def test_upload_valid_docx():
    docx_bytes = create_sample_docx_bytes()
    response = client.post(
        "/api/documents/upload",
        files={"file": ("code_of_conduct.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["message"] == "Policy uploaded and processed successfully."
    assert data["document"]["filename"] == "code_of_conduct.docx"
    assert data["document"]["file_type"] == "docx"
    assert data["document"]["chunk_count"] >= 2


def test_upload_unsupported_file():
    txt_bytes = b"Hello, this is a plain text file."
    response = client.post(
        "/api/documents/upload",
        files={"file": ("test.txt", txt_bytes, "text/plain")}
    )
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]


def test_upload_empty_file():
    empty_bytes = b""
    response = client.post(
        "/api/documents/upload",
        files={"file": ("empty_policy.pdf", empty_bytes, "application/pdf")}
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()
